#!/usr/bin/env python3
"""Сборка отчёта Q2 2026 с фирменным стилем QazCloud (синий + Graphik LCG)."""

import importlib.util
import shutil
import subprocess
import sys
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "report-q2-2026" / "screenshots"
FONTS = ROOT / "docs" / "report-q2-2026" / "fonts"
OUT = ROOT / "docs" / "report-q2-2026" / "Отчет_Лаборатория_Q2_2026.docx"
DOWNLOADS_DOCX = Path.home() / "Downloads" / "Отчет_Лаборатория_Q2_2026.docx"
PDF_SCRIPT = ROOT / "scripts" / "build-lab-report-pdf.py"

FONT_FAMILY = "Graphik LCG"
FONT_FILES = {
    "regular": FONTS / "GraphikLCG-Regular.ttf",
    "semibold": FONTS / "GraphikLCG-Semibold.ttf",
    "bold": FONTS / "GraphikLCG-Bold.ttf",
}

# QazCloud brand palette
BLUE_DEEP = RGBColor(0x24, 0x38, 0xC7)
BLUE_MID = RGBColor(0x64, 0x78, 0xF5)
BLUE_LIGHT = RGBColor(0x8B, 0xAD, 0xFF)
BLUE_SURFACE = RGBColor(0x4B, 0x62, 0xEA)
TEXT_BODY = RGBColor(0x1E, 0x29, 0x3B)
TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEX_BLUE_DEEP = "2438C7"
HEX_BLUE_MID = "6478F5"
HEX_ROW_ALT = "F0F4FF"

CAPTIONS: dict[str, str] = {}


def load_report_data():
    path = ROOT / "scripts" / "report-lab-data.py"
    spec = importlib.util.spec_from_file_location("report_lab_data", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.SECTIONS, mod.CAPTIONS


def set_run_font(
    run,
    *,
    size=11,
    bold=False,
    italic=False,
    color=TEXT_BODY,
    weight="regular",
):
    run.font.name = FONT_FAMILY
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

    r_pr = run._element.get_or_add_rPr()
    for child in list(r_pr.findall(qn("w:rFonts"))):
        r_pr.remove(child)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_FAMILY)
    r_fonts.set(qn("w:hAnsi"), FONT_FAMILY)
    r_fonts.set(qn("w:cs"), FONT_FAMILY)
    r_fonts.set(qn("w:eastAsia"), FONT_FAMILY)
    if weight == "semibold":
        r_fonts.set(qn("w:asciiTheme"), "")
    r_pr.insert(0, r_fonts)


def shade_paragraph(paragraph, fill_hex):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    p_pr.append(shd)


def shade_cell(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def apply_document_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_FAMILY
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT_BODY

    for level, size in ((1, 16), (2, 13), (3, 12)):
        style_name = f"Heading {level}"
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = FONT_FAMILY
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = BLUE_DEEP

    for style_name in ("List Bullet", "List Number"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = FONT_FAMILY
            style.font.size = Pt(11)
            style.font.color.rgb = TEXT_BODY


def add_title_page(doc):
    band = doc.add_paragraph()
    band.alignment = WD_ALIGN_PARAGRAPH.LEFT
    shade_paragraph(band, HEX_BLUE_DEEP)
    band.paragraph_format.space_before = Pt(0)
    band.paragraph_format.space_after = Pt(0)
    band.paragraph_format.left_indent = Inches(-1.0)
    band.paragraph_format.right_indent = Inches(-1.0)

    spacer = band.add_run("\n\n")
    set_run_font(spacer, size=6, color=WHITE)

    title = band.add_run("AI-HUB Enterprise Platform")
    set_run_font(title, size=22, bold=True, color=WHITE, weight="bold")

    band.add_run("\n")
    subtitle = band.add_run("UX/UI-дизайн интерактивного прототипа: «Лаборатория»")
    set_run_font(subtitle, size=15, bold=True, color=BLUE_LIGHT, weight="semibold")

    band.add_run("\n\n")
    meta = band.add_run("Отчёт UX/UI-дизайнера за II квартал 2026 года")
    set_run_font(meta, size=12, color=WHITE)

    band.add_run("\n\n")
    set_run_font(band.add_run(), size=4, color=WHITE)

    accent = doc.add_paragraph()
    shade_paragraph(accent, HEX_BLUE_MID)
    accent.paragraph_format.space_before = Pt(0)
    accent.paragraph_format.space_after = Pt(18)
    accent.paragraph_format.left_indent = Inches(-1.0)
    accent.paragraph_format.right_indent = Inches(-1.0)
    set_run_font(accent.add_run(" "), size=3, color=WHITE)

    doc.add_paragraph()


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(
            run,
            size={1: 16, 2: 13, 3: 12}.get(level, 12),
            bold=True,
            color=BLUE_DEEP,
            weight="semibold",
        )
    return h


def add_para(doc, text, bold=False, color=TEXT_BODY, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, weight="bold" if bold else "regular")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run, size=11, color=TEXT_BODY)
        p.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(header)
        set_run_font(run, size=10, bold=True, color=WHITE, weight="semibold")
        shade_cell(cell, HEX_BLUE_DEEP)
        set_cell_margins(cell)

    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cell = cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10, color=TEXT_BODY)
            if ri % 2 == 1:
                shade_cell(cell, HEX_ROW_ALT)
            set_cell_margins(cell)

    doc.add_paragraph()


def add_screenshot(doc, shot_id: str):
    filename = shot_id if shot_id.endswith(".png") else f"{shot_id}.png"
    path = SHOTS / filename
    if not path.exists():
        add_para(doc, f"[Скриншот не найден: {filename}]", bold=True, color=BLUE_DEEP)
        return

    doc.add_paragraph()
    frame = doc.add_paragraph()
    frame.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade_paragraph(frame, HEX_ROW_ALT)
    frame.paragraph_format.left_indent = Inches(0.15)
    frame.paragraph_format.right_indent = Inches(0.15)
    frame.paragraph_format.space_before = Pt(8)
    frame.paragraph_format.space_after = Pt(8)

    run = frame.add_run()
    run.add_picture(str(path), width=Inches(6.2))

    cap = doc.add_paragraph(CAPTIONS.get(filename.removesuffix(".png"), filename))
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        set_run_font(run, size=9, italic=True, color=TEXT_MUTED)
    doc.add_paragraph()


def embed_fonts(docx_path: Path):
    """Встраивает Graphik LCG в docx для корректного отображения без установки шрифта."""
    font_entries = [
        ("regular", FONT_FILES["regular"], "embedRegular"),
        ("bold", FONT_FILES["bold"], "embedBold"),
    ]

    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    ET.register_namespace("w", ns["w"])
    ET.register_namespace("r", ns["r"])

    with zipfile.ZipFile(docx_path, "r") as zin:
        archive = {name: zin.read(name) for name in zin.namelist()}

    rels_path = "word/_rels/fontTable.xml.rels"
    font_table_path = "word/fontTable.xml"
    content_types_path = "[Content_Types].xml"

    if rels_path not in archive:
        archive[rels_path] = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            b"</Relationships>"
        )
    if font_table_path not in archive:
        archive[font_table_path] = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<w:fonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            b' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            b"</w:fonts>"
        )

    rels_root = ET.fromstring(archive[rels_path])
    rels_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    existing_rels = rels_root.findall(f"{{{rels_ns}}}Relationship")
    next_rid = len(existing_rels) + 1

    font_root = ET.fromstring(archive[font_table_path])
    for child in list(font_root):
        if child.tag == f"{{{ns['w']}}}font" and child.get(f"{{{ns['w']}}}name") == FONT_FAMILY:
            font_root.remove(child)

    font_el = ET.SubElement(font_root, f"{{{ns['w']}}}font")
    font_el.set(f"{{{ns['w']}}}name", FONT_FAMILY)
    ET.SubElement(font_el, f"{{{ns['w']}}}panose1", {f"{{{ns['w']}}}val": "020B0604030504030204"})
    ET.SubElement(font_el, f"{{{ns['w']}}}charset", {f"{{{ns['w']}}}val": "CC"})
    ET.SubElement(font_el, f"{{{ns['w']}}}family", {f"{{{ns['w']}}}val": "swiss"})

    ct_root = ET.fromstring(archive[content_types_path])
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"

    for _, font_path, embed_tag in font_entries:
        if not font_path.exists():
            continue

        part_name = f"/word/fonts/{font_path.name}"
        archive[f"word/fonts/{font_path.name}"] = font_path.read_bytes()

        override_exists = any(
            el.get("PartName") == part_name for el in ct_root.findall(f"{{{ct_ns}}}Override")
        )
        if not override_exists:
            ET.SubElement(
                ct_root,
                f"{{{ct_ns}}}Override",
                {"PartName": part_name, "ContentType": "application/x-fontdata"},
            )

        rid = f"rId{next_rid}"
        next_rid += 1
        ET.SubElement(
            rels_root,
            f"{{{rels_ns}}}Relationship",
            {
                "Id": rid,
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
                "Target": f"fonts/{font_path.name}",
            },
        )
        embed_el = ET.SubElement(font_el, f"{{{ns['w']}}}{embed_tag}")
        embed_el.set(f"{{{ns['r']}}}id", rid)
        embed_el.set(f"{{{ns['w']}}}fontKey", "{" + str(uuid.uuid4()).upper() + "}")

    archive[rels_path] = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)
    archive[font_table_path] = ET.tostring(font_root, encoding="utf-8", xml_declaration=True)
    archive[content_types_path] = ET.tostring(ct_root, encoding="utf-8", xml_declaration=True)

    with zipfile.ZipFile(docx_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in archive.items():
            zout.writestr(name, data)


def render_sections(doc, sections):
    for section in sections:
        add_heading(doc, section["title"], 1)
        for paragraph in section.get("paragraphs", []):
            add_para(doc, paragraph)
        if "bullets" in section:
            add_bullets(doc, section["bullets"])
        if "table" in section:
            table = section["table"]
            add_table(doc, table["headers"], table["rows"])
        for shot_id in section.get("screenshots", []):
            add_screenshot(doc, shot_id)


def build():
    global CAPTIONS
    OUT.parent.mkdir(parents=True, exist_ok=True)
    FONTS.mkdir(parents=True, exist_ok=True)

    sections, CAPTIONS = load_report_data()

    doc = Document()
    apply_document_styles(doc)
    add_title_page(doc)
    render_sections(doc, sections)

    doc.save(OUT)
    embed_fonts(OUT)
    shutil.copy2(OUT, DOWNLOADS_DOCX)
    print(f"Saved DOCX: {OUT}")
    print(f"Copy DOCX:  {DOWNLOADS_DOCX}")

    subprocess.run([sys.executable, str(PDF_SCRIPT)], check=True)


if __name__ == "__main__":
    build()
